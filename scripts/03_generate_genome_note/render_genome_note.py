#!/usr/bin/env python3
"""
Render a genome note JATS XML file to HTML or Word (.docx).

Usage:
  python3 render_genome_note.py OG38_genome_note.xml          # → HTML
  python3 render_genome_note.py OG38_genome_note.xml --docx   # → Word doc
  python3 render_genome_note.py OG38_genome_note.xml -o out.html

For --docx, python-docx must be installed:
  uv run --with python-docx render_genome_note.py OG38_genome_note.xml --docx
"""

import sys
import re
import base64
import argparse
from pathlib import Path
import xml.etree.ElementTree as ET

try:
    from docx import Document as DocxDocument
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    _DOCX_AVAILABLE = True
except ImportError:
    _DOCX_AVAILABLE = False

try:
    from PIL import Image as _PILImage
    import io as _io
    _PIL_AVAILABLE = True
except ImportError:
    _PIL_AVAILABLE = False

NS = {
    "xlink": "http://www.w3.org/1999/xlink",
    "mml":   "http://www.w3.org/1998/Math/MathML",
}

# ── XML helpers ────────────────────────────────────────────────────────────────

def text(el) -> str:
    """Recursively collect all text from an element."""
    if el is None:
        return ""
    parts = [el.text or ""]
    for child in el:
        parts.append(inline(child))
        parts.append(child.tail or "")
    return "".join(parts)


def inline(el) -> str:
    """Render an element and its children as inline HTML."""
    if el is None:
        return ""
    tag = el.tag.split("}")[-1] if "}" in el.tag else el.tag
    inner = (el.text or "") + "".join(inline(c) + (c.tail or "") for c in el)

    if tag in ("italic", "i"):
        return f"<em>{inner}</em>"
    if tag in ("bold", "b"):
        return f"<strong>{inner}</strong>"
    if tag == "db-val":
        return f'<strong class="db-val">{inner}</strong>'
    if tag == "sup":
        return f"<sup>{inner}</sup>"
    if tag == "sub":
        return f"<sub>{inner}</sub>"
    if tag == "break":
        return "<br>"
    if tag == "ext-link":
        href = el.get("{http://www.w3.org/1999/xlink}href") or el.get("xlink:href") or "#"
        return f'<a href="{href}" target="_blank">{inner}</a>'
    if tag == "xref":
        rid  = el.get("rid", "")
        kind = el.get("ref-type", "")
        if kind == "fig":
            return f'<a href="#{rid}" class="xref">{inner}</a>'
        if kind == "table":
            return f'<a href="#{rid}" class="xref">{inner}</a>'
        if kind == "bibr":
            return f'<a href="#{rid}" class="xref">{inner}</a>'
        return f'<a href="#{rid}" class="xref">{inner}</a>'
    if tag == "fn":
        return f'<span class="fn-inline">{inner}</span>'
    # fallback — just return the text
    return inner


def para(el) -> str:
    """Render a <p> element."""
    if el is None:
        return ""
    inner = (el.text or "") + "".join(inline(c) + (c.tail or "") for c in el)
    return f"<p>{inner}</p>"


# ── Section renderers ──────────────────────────────────────────────────────────

def render_contrib_group(root: ET.Element) -> str:
    parts = []
    for contrib in root.findall(".//{*}contrib[@contrib-type='author']"):
        name_el = contrib.find("{*}name")
        collab_el = contrib.find("{*}collab")
        aff_refs = [text(x) for x in contrib.findall("{*}xref[@ref-type='aff']") if text(x)]
        sup_html = f"<sup>{','.join(aff_refs)}</sup>" if aff_refs else ""
        if name_el is not None:
            given = text(name_el.find("{*}given-names"))
            surname = text(name_el.find("{*}surname"))
            parts.append(f"{given} {surname}{sup_html}")
        elif collab_el is not None:
            parts.append(f"<em>{text(collab_el)}</em>")
    return ", ".join(parts)


def render_affiliations(root: ET.Element) -> str:
    affs = root.findall(".//{*}aff")
    if not affs:
        return ""
    items = []
    for aff in affs:
        label = text(aff.find("{*}label"))
        aff_text = (aff.text or "") + "".join(inline(c) + (c.tail or "") for c in aff)
        aff_text = re.sub(r"^\s*\d+\s*", "", aff_text).strip()
        if label:
            items.append(f"<sup>{label}</sup> {aff_text}")
        else:
            items.append(aff_text)
    return "<p class='affiliations'>" + "<br>".join(items) + "</p>"


def render_abstract(root: ET.Element) -> str:
    abstract = root.find(".//{*}abstract")
    if abstract is None:
        return ""
    paras = "".join(para(p) for p in abstract.findall("{*}p"))
    return f"<h2>Abstract</h2>{paras}"


def render_keywords(root: ET.Element) -> str:
    kwds = root.findall(".//{*}kwd")
    if not kwds:
        return ""
    items = [f"<span class='kwd'>{inline(k)}</span>" for k in kwds]
    return f"<strong>Keywords:</strong> {''.join(items)}"


def render_table(tbl_wrap: ET.Element) -> str:
    tid   = tbl_wrap.get("id", "")
    label = text(tbl_wrap.find("{*}label"))
    cap   = tbl_wrap.find(".//{*}caption")
    title_el = cap.find("{*}title") if cap is not None else None
    cap_text = (inline(title_el) if title_el is not None else "") + \
               "".join(para(p) for p in (cap.findall("{*}p") if cap is not None else []))
    cap_text = cap_text.strip()

    tbl = tbl_wrap.find(".//{*}table")
    if tbl is None:
        return ""

    rows_html = []
    for row in tbl.findall(".//{*}tr"):
        cells = []
        for cell in row:
            cell_tag = cell.tag.split("}")[-1]
            inner = (cell.text or "") + "".join(inline(c) + (c.tail or "") for c in cell)
            colspan = cell.get("colspan", "1")
            rowspan = cell.get("rowspan", "1")
            align   = cell.get("align", "left")
            span_attrs = ""
            if colspan != "1": span_attrs += f' colspan="{colspan}"'
            if rowspan != "1": span_attrs += f' rowspan="{rowspan}"'
            tag = "th" if cell_tag == "th" else "td"
            cells.append(f'<{tag}{span_attrs} style="text-align:{align}">{inner}</{tag}>')
        rows_html.append("<tr>" + "".join(cells) + "</tr>")

    # footnotes
    feet = []
    for fn in tbl_wrap.findall(".//{*}fn"):
        for p_el in fn.findall("{*}p"):
            feet.append(f"<p class='table-fn'>{para(p_el)}</p>")

    return f"""
    <div class="table-wrap" id="{tid}">
      <p class="table-label">{label} {cap_text}</p>
      <table>{''.join(rows_html)}</table>
      {''.join(feet)}
    </div>"""


def render_figure(fig: ET.Element) -> str:
    fid   = fig.get("id", "")
    label = text(fig.find("{*}label"))
    cap   = fig.find("{*}caption")
    title_el = cap.find("{*}title") if cap is not None else None
    cap_title = inline(title_el) if title_el is not None else ""
    cap_paras = "".join(para(p) for p in (cap.findall("{*}p") if cap is not None else []))

    graphic = fig.find("{*}graphic")
    src = ""
    if graphic is not None:
        src = graphic.get("{http://www.w3.org/1999/xlink}href") or \
              graphic.get("xlink:href") or ""

    is_missing = not src or src.startswith("[MISSING:") or src.startswith("PLACEHOLDER")
    if is_missing:
        img_html = f'<div class="placeholder-img">{src if src else "No path provided"}</div>'
    else:
        embedded = _embed_image(src)
        img_html = f'<img src="{embedded}" alt="{label}">'

    return f"""
    <figure id="{fid}">
      {img_html}
      <figcaption><strong>{label}</strong> {cap_title}{cap_paras}</figcaption>
    </figure>"""


def render_sec(sec: ET.Element, depth: int = 2) -> str:
    title_el = sec.find("{*}title")
    title_text = text(title_el) if title_el is not None else ""
    htag = f"h{min(depth, 4)}"

    parts = [f"<{htag}>{title_text}</{htag}>"]

    for child in sec:
        ctag = child.tag.split("}")[-1]
        if ctag == "title":
            continue
        elif ctag == "p":
            parts.append(para(child))
        elif ctag == "sec":
            parts.append(render_sec(child, depth + 1))
        elif ctag == "table-wrap":
            parts.append(render_table(child))
        elif ctag == "fig":
            parts.append(render_figure(child))
        elif ctag in ("ref-list",):
            parts.append(render_refs(child))
        else:
            # generic: collect paragraphs inside
            for p_el in child.findall(".//{*}p"):
                parts.append(para(p_el))

    return "<section>" + "".join(parts) + "</section>"


def render_refs(ref_list: ET.Element) -> str:
    items = []
    for ref in ref_list.findall("{*}ref"):
        rid = ref.get("id", "")
        ec = ref.find(".//{*}element-citation")
        if ec is None:
            ec = ref.find(".//{*}mixed-citation")
        if ec is None:
            continue

        authors = []
        for pg in ec.findall("{*}person-group"):
            for name in pg.findall("{*}name"):
                s = text(name.find("{*}surname"))
                g = text(name.find("{*}given-names"))
                authors.append(f"{s} {g[0]}" if g else s)
            for collab in pg.findall("{*}collab"):
                ct = text(collab)
                if ct:
                    authors.append(ct)
        author_str = ", ".join(a for a in authors if a)

        title  = text(ec.find("{*}article-title"))
        source = text(ec.find("{*}source"))
        year   = text(ec.find("{*}year"))
        vol    = text(ec.find("{*}volume"))
        fp     = text(ec.find("{*}fpage"))
        lp     = text(ec.find("{*}lpage"))
        doi_el = ec.find("{*}pub-id[@pub-id-type='doi']")
        doi    = text(doi_el) if doi_el is not None else ""

        # Author + year
        bib = f"{author_str} ({year})." if author_str else f"({year})."
        # Title (italic)
        if title:
            bib += f" <em>{title}</em>."
        # Source + volume + pages (only if source exists)
        if source:
            src = source
            if vol:
                src += f", {vol}"
            if fp:
                src += f":{fp}–{lp}" if lp else f":{fp}"
            bib += f" {src}."
        # DOI link
        if doi:
            bib += f' <a href="https://doi.org/{doi}" target="_blank">doi:{doi}</a>.'

        items.append(f'<p id="{rid}" class="ref">{bib}</p>')

    return "".join(items)


def render_back(back: ET.Element) -> str:
    if back is None:
        return ""
    parts = []
    for sec in back.findall("{*}sec"):
        parts.append(render_sec(sec, depth=2))
    return "".join(parts)


# ── CSS ────────────────────────────────────────────────────────────────────────

CSS = """
@import url('https://fonts.googleapis.com/css2?family=Lora:ital,wght@0,400;0,500;0,700;1,400;1,500;1,700&display=swap');

* { box-sizing: border-box; margin: 0; padding: 0; }

body {
  font-family: 'Lora', Georgia, 'Times New Roman', serif;
  font-size: 10.5pt;
  line-height: 1.65;
  color: #1a1a1a;
  background: #d4d4d4;
  padding: 28px 16px;
}

/* ── Page shell ── */
.page {
  max-width: 760px;
  margin: 0 auto;
  background: #fff;
  box-shadow: 0 2px 12px rgba(0,0,0,.22);
}

/* ── Article header ── */
.article-header {
  padding: 20px 48px 18px;
  border-bottom: 1px solid #ddd;
}
.article-type {
  display: block;
  font-size: 8pt;
  font-weight: 700;
  color: #007475;
  margin-bottom: 8px;
}
h1.article-title {
  font-size: 18pt;
  font-weight: 700;
  line-height: 1.22;
  color: #007475;
  margin-bottom: 14px;
}
.authors {
  font-size: 10pt;
  color: #222;
  margin-bottom: 8px;
  line-height: 1.75;
}
.authors sup { font-size: 6.5pt; color: #007475; vertical-align: super; }
.affiliations {
  font-size: 8.5pt;
  color: #555;
  line-height: 1.6;
  margin-bottom: 10px;
}
.affiliations sup { font-size: 6pt; color: #007475; vertical-align: super; }
.article-meta {
  font-size: 8pt;
  color: #777;
  padding: 8px 0 0;
  border-top: 1px solid #ddd;
  margin-top: 10px;
}

/* ── Abstract & keywords — no gray box, inline with page ── */
.abstract-wrap {
  padding: 22px 48px 6px;
}
.abstract-wrap h2 {
  font-size: 10.5pt;
  font-weight: 700;
  color: #007475;
  border: none;
  margin: 0 0 8px;
  padding: 0;
}
.abstract-wrap p {
  font-size: 10pt;
  line-height: 1.65;
  text-align: justify;
  margin-bottom: 0;
}
.keywords-wrap {
  padding: 8px 48px 22px;
  border-bottom: 1px solid #ccc;
  font-size: 10pt;
  color: #333;
}
.keywords-wrap strong { font-weight: 700; color: #007475; }
.kwd { display: inline; font-style: italic; color: #333; }
.kwd + .kwd::before { content: ", "; font-style: normal; }

/* ── Body ── */
.article-body { padding: 26px 48px 40px; }

/* Top-level sections get a thin rule above them to visually separate content */
.article-body > section {
  border-top: 1px solid #ddd;
  padding-top: 18px;
  margin-top: 0;
  margin-bottom: 28px;
}
.article-body > section:first-child {
  border-top: none;
  padding-top: 0;
}

section { margin-bottom: 4px; }

/* Level-1 body sections (Background, Methods, etc.): bold, dark */
h2 {
  font-size: 11pt;
  font-weight: 700;
  color: #111;
  margin: 0 0 10px;
}

/* Level-2 subsections (DNA extraction, Hi-C, etc.): teal, normal weight */
h3 {
  font-size: 10.5pt;
  font-weight: 400;
  color: #007475;
  margin: 18px 0 5px;
}

/* Level-3 sub-subsections: bold italic, dark */
h4 {
  font-size: 10.5pt;
  font-weight: 700;
  font-style: italic;
  color: #111;
  margin: 14px 0 4px;
}

p { margin-bottom: 10px; text-align: justify; }
em { font-style: italic; }
strong { font-weight: 700; }
sup, sub { font-size: 6.5pt; line-height: 0; position: relative; vertical-align: baseline; }
sup { top: -0.5em; }
sub { bottom: -0.25em; }

a { color: #007475; text-decoration: underline; }
a:hover { text-decoration: none; }

/* ── Tables ── */
.table-wrap { margin: 24px 0; overflow-x: auto; }
.table-label {
  font-size: 9.5pt;
  font-weight: 700;
  color: #007475;
  margin-bottom: 3px;
}
.table-caption {
  font-size: 9pt;
  color: #333;
  margin-bottom: 8px;
  line-height: 1.5;
}
table {
  border-collapse: collapse;
  width: 100%;
  font-size: 9pt;
}
thead tr {
  background: #d9eff0;
  border-top: 2px solid #007475;
  border-bottom: 1px solid #007475;
}
tbody tr:last-child { border-bottom: 2px solid #007475; }
th {
  color: #111;
  font-weight: 700;
  text-align: center;
  padding: 5px 10px;
  vertical-align: bottom;
}
td {
  padding: 4px 10px;
  vertical-align: top;
  line-height: 1.45;
  border: none;
}
td:first-child { font-weight: 700; text-align: left; }
td:not(:first-child) { text-align: center; }
tr:nth-child(even) td { background: #f0fafb; }
.table-fn { font-size: 8pt; color: #555; margin-top: 5px; }

/* ── Figures ── */
figure { margin: 26px 0; }
figure img {
  display: block;
  max-width: 85%;
  max-height: 500px;
  width: auto;
  height: auto;
  margin: 0 auto;
}
.placeholder-img {
  background: #f5f5f5;
  border: 1px dashed #bbb;
  padding: 36px 20px;
  color: #888;
  font-size: 9pt;
  text-align: center;
  font-style: italic;
}
figcaption {
  font-size: 9pt;
  color: #222;
  margin-top: 9px;
  line-height: 1.55;
  text-align: justify;
}
figcaption strong { font-weight: 700; color: #007475; }

/* ── References ── */
.ref {
  font-size: 9pt;
  color: #222;
  line-height: 1.6;
  margin-bottom: 8px;
  padding-left: 2em;
  text-indent: -2em;
  text-align: left;
}
.ref em { font-style: italic; }
.ref a { color: #007475; }

a.xref, .xref { color: #007475; text-decoration: none; }
a.xref:hover, .xref:hover { text-decoration: underline; }

/* ── Back matter sections ── */
.back-section { margin-top: 20px; font-size: 10pt; }

/* ── Print button ── */
.print-bar {
  position: fixed;
  top: 14px;
  right: 18px;
  z-index: 99;
}
.print-bar button {
  background: #007475;
  color: #fff;
  border: none;
  padding: 7px 16px;
  border-radius: 3px;
  cursor: pointer;
  font-size: 10pt;
  font-family: inherit;
  box-shadow: 0 2px 6px rgba(0,0,0,.25);
}
.print-bar button:hover { background: #005c5d; }

/* ── DB highlight (review mode) ── */
.db-val {
  color: #0a56a0;
  background: #ddeeff;
  border-radius: 2px;
  padding: 0 2px;
  font-weight: 600;
}

/* ── Review banner ── */
.review-banner {
  background: #fff8e1;
  border-bottom: 2px solid #f9a825;
  padding: 10px 48px;
  font-size: 9.5pt;
  color: #5d4037;
}
.review-banner strong { color: #0a56a0; background: #ddeeff; border-radius: 2px; padding: 0 3px; }

@media print {
  body { background: #fff; padding: 0; }
  .page { box-shadow: none; }
  .print-bar { display: none; }
  .review-banner { display: none; }
  h2 { page-break-after: avoid; }
  figure, .table-wrap { page-break-inside: avoid; }
}
"""


_XML_DIR: Path = Path(".")


def _embed_image(src: str) -> str:
    """Return a data URI for src if the file exists, otherwise return src unchanged."""
    img_path = (_XML_DIR / src).resolve()
    if not img_path.exists():
        return src
    suffix = img_path.suffix.lower().lstrip(".")
    mime = {"jpg": "jpeg", "jpeg": "jpeg", "png": "png", "gif": "gif",
            "svg": "svg+xml", "webp": "webp"}.get(suffix, suffix)
    data = base64.b64encode(img_path.read_bytes()).decode()
    return f"data:image/{mime};base64,{data}"


# ── Main renderer ──────────────────────────────────────────────────────────────

def render(xml_path: Path) -> str:
    global _XML_DIR
    _XML_DIR = xml_path.parent
    tree = ET.parse(xml_path)
    root = tree.getroot()

    # Namespace-aware tag stripping already handled by split("}")[-1] in helpers.
    # Register namespaces so ET doesn't rewrite them.
    ET.register_namespace("xlink", "http://www.w3.org/1999/xlink")

    front  = root.find("{*}front")
    body   = root.find("{*}body")
    back_el = root.find("{*}back")

    # ── Front matter ──
    article_meta = root.find(".//{*}article-meta")
    title_el = root.find(".//{*}article-title")
    article_title = (title_el.text or "") + "".join(inline(c) + (c.tail or "") for c in title_el) if title_el is not None else "Genome Note"

    journal_el = root.find(".//{*}journal-title")
    journal_title = text(journal_el) if journal_el is not None else "Wellcome Open Research"

    authors_html = render_contrib_group(root)
    affs_html    = render_affiliations(root)
    abstract_html = render_abstract(root)
    kwd_html     = render_keywords(root)

    pub_date = root.find(".//{*}pub-date[@pub-type='epub']")
    pub_str = ""
    if pub_date is not None:
        day   = text(pub_date.find("{*}day"))
        month = text(pub_date.find("{*}month"))
        year  = text(pub_date.find("{*}year"))
        from calendar import month_abbr
        try:
            month_name = month_abbr[int(month)]
        except (ValueError, IndexError):
            month_name = month
        pub_str = f"{day} {month_name} {year}" if all([day, month_name, year]) else ""

    doi_el = root.find(".//{*}article-id[@pub-id-type='doi']")
    doi_str = text(doi_el) if doi_el is not None else ""

    # ── Body sections ──
    body_html = ""
    if body is not None:
        for child in body:
            ctag = child.tag.split("}")[-1]
            if ctag == "sec":
                body_html += render_sec(child, depth=2)

    # ── Back matter ──
    back_html = render_back(back_el) if back_el is not None else ""

    # ── Assemble ──
    meta_parts = []
    if pub_str and "[MISSING" not in pub_str and "PLACEHOLDER" not in pub_str:
        meta_parts.append(f"Published: {pub_str}")
    if doi_str and "[MISSING" not in doi_str and "PLACEHOLDER" not in doi_str:
        meta_parts.append(f'DOI: <a href="https://doi.org/{doi_str}">{doi_str}</a>')
    meta_line = " &nbsp;&bull;&nbsp; ".join(meta_parts)

    plain_title = re.sub('<[^>]+>', '', article_title)

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{plain_title}</title>
  <style>{CSS}</style>
</head>
<body>
  <div class="print-bar">
    <button onclick="window.print()">&#128438; Save as PDF</button>
  </div>
  <div class="page">

    <div class="article-header">
      <div class="article-type">Data Note</div>
      <h1 class="article-title">{article_title}</h1>
      <p class="authors">{authors_html}</p>
      {affs_html}
      {"<div class='article-meta'>" + meta_line + "</div>" if meta_line else ""}
    </div>

    <div class="abstract-wrap">
      {abstract_html}
    </div>
    <div class="keywords-wrap">
      {kwd_html}
    </div>

    <div class="article-body">
      {body_html}
      {back_html}
      <p style="font-size:7.5pt;color:#aaa;text-align:right;margin-top:30px;border-top:1px solid #eee;padding-top:8px">
        Generated by OceanOmics genome notes automation pipeline &mdash; preview only
      </p>
    </div>

  </div>
</body>
</html>"""


# ── Word doc renderer ─────────────────────────────────────────────────────────

def _ws(s: str) -> str:
    """Collapse XML indentation whitespace (newlines + runs of spaces) to one space."""
    return re.sub(r'\s+', ' ', s or '')


def _plain(el) -> str:
    """Strip all XML tags, normalize whitespace, and return plain text."""
    if el is None:
        return ""
    parts = [el.text or ""]
    for child in el:
        parts.append(_plain(child))
        parts.append(child.tail or "")
    return _ws("".join(parts)).strip()


# Bookmark counter — must be unique across the whole document
_bmark_counter = 0


def _next_bmark_id() -> int:
    global _bmark_counter
    _bmark_counter += 1
    return _bmark_counter


def _add_bookmark(para, name: str):
    """Wrap paragraph content with a named Word bookmark (jump target)."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    bid = str(_next_bmark_id())
    p = para._p
    bstart = OxmlElement("w:bookmarkStart")
    bstart.set(qn("w:id"), bid)
    bstart.set(qn("w:name"), name)
    bend = OxmlElement("w:bookmarkEnd")
    bend.set(qn("w:id"), bid)
    pPr = p.find(qn("w:pPr"))
    idx = list(p).index(pPr) + 1 if pPr is not None else 0
    p.insert(idx, bstart)
    p.append(bend)


def _add_external_link(para, display_text: str, url: str, bold=False, italic=False):
    """Add a blue underlined external URL hyperlink to a paragraph."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    r_id = para.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    if bold:
        rPr.append(OxmlElement("w:b"))
    if italic:
        rPr.append(OxmlElement("w:i"))
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = display_text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    hl.append(r)
    para._p.append(hl)


def _add_internal_link(para, display_text: str, anchor: str, bold=False, italic=False):
    """Add a blue underlined internal hyperlink (bookmark anchor) to a paragraph."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("w:anchor"), anchor)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    if bold:
        b = OxmlElement("w:b")
        rPr.append(b)
    if italic:
        i_el = OxmlElement("w:i")
        rPr.append(i_el)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = display_text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    hl.append(r)
    para._p.append(hl)


def _add_run(para, el, bold=False, italic=False):
    """Recursively add runs from an XML element, collapsing XML indentation whitespace.
    xref[@ref-type='bibr'] elements are rendered as internal hyperlinks."""
    if el.text:
        txt = _ws(el.text)
        if txt:
            run = para.add_run(txt)
            run.bold = bold
            run.italic = italic
    for child in el:
        ctag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        child_italic = italic or ctag in ("italic", "i")
        child_bold   = bold   or ctag in ("bold", "b")

        if ctag == "xref" and child.get("ref-type") == "bibr":
            rid = child.get("rid", "")
            link_txt = _plain(child)
            if rid and link_txt:
                _add_internal_link(para, link_txt, rid, bold=bold, italic=italic)
            elif link_txt:
                run = para.add_run(link_txt)
                run.bold = bold
                run.italic = italic
        elif ctag == "ext-link":
            url = child.get("{http://www.w3.org/1999/xlink}href") or child.get("xlink:href") or ""
            link_txt = _plain(child)
            if url and link_txt:
                _add_external_link(para, link_txt, url, bold=bold, italic=italic)
            elif link_txt:
                run = para.add_run(link_txt)
                run.bold = bold
                run.italic = italic
        else:
            _add_run(para, child, bold=child_bold, italic=child_italic)

        if child.tail:
            txt = _ws(child.tail)
            if txt:
                run = para.add_run(txt)
                run.bold = bold
                run.italic = italic


def _set_table_borders(tbl):
    """Force visible single-line black borders on all table cells."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tblPr = tbl._tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl._tbl.insert(0, tblPr)
    existing = tblPr.find(qn("w:tblBorders"))
    if existing is not None:
        tblPr.remove(existing)
    tblBorders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")   # ½ pt
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")
        tblBorders.append(border)
    tblPr.append(tblBorders)


def _style_exists(doc, name: str) -> bool:
    return any(s.name == name for s in doc.styles)


def render_docx(xml_path: Path) -> "DocxDocument":
    """Build a python-docx Document from a JATS XML genome note.

    Opens GN_Structure_template.docx (if present) as the base so all
    fonts, heading styles, and page layout are inherited from the template.
    """
    global _bmark_counter
    _bmark_counter = 0  # reset so IDs are clean for each document

    tree = ET.parse(xml_path)
    root = tree.getroot()

    # ── Open template to inherit styles / layout ──────────────────────────────
    template_path = xml_path.parent / "GN_Structure_template.docx"
    if template_path.exists():
        doc = DocxDocument(str(template_path))
        # Clear all body content, keeping styles and section properties
        from docx.oxml.ns import qn as _qn
        body = doc.element.body
        for child in list(body):
            if child.tag != _qn("w:sectPr"):
                body.remove(child)
    else:
        doc = DocxDocument()
        doc.styles["Normal"].font.name = "Calibri"
        doc.styles["Normal"].font.size = Pt(11)

    # ── Title ─────────────────────────────────────────────────────────────────
    title_el = root.find(".//{*}article-title")
    if title_el is not None:
        t = doc.add_heading(level=0)
        t.clear()
        _add_run(t, title_el)

    # ── Authors (with superscript affiliation numbers) ────────────────────────
    authors_html = render_contrib_group(root)
    if authors_html:
        p = doc.add_paragraph()
        for chunk in re.split(r'(<sup>[^<]*</sup>|<[^>]+>)', authors_html):
            sup_m = re.match(r'<sup>([^<]*)</sup>', chunk)
            tag_m = re.match(r'<[^>]+>', chunk)
            if sup_m:
                run = p.add_run(sup_m.group(1))
                run.font.superscript = True
            elif tag_m:
                pass  # skip other HTML tags (em, strong, etc.)
            elif chunk:
                p.add_run(_ws(chunk))

    # ── Affiliations (deduplicated by text) ───────────────────────────────────
    seen_affs: dict[str, str] = {}  # aff_text → first label seen
    for aff in root.findall(".//{*}aff"):
        label = _plain(aff.find("{*}label"))
        aff_txt = _ws((aff.text or "") + "".join(
            (_plain(c) + _ws(c.tail or "")) for c in aff
        )).strip()
        aff_txt = re.sub(r"^\s*\d+\s*", "", aff_txt).strip()
        if aff_txt and aff_txt not in seen_affs:
            seen_affs[aff_txt] = label
    for aff_txt, label in seen_affs.items():
        p = doc.add_paragraph()
        if label:
            p.add_run(f"{label} ").bold = True
        p.add_run(aff_txt)

    # ── Abstract ──────────────────────────────────────────────────────────────
    abstract = root.find(".//{*}abstract")
    if abstract is not None:
        doc.add_heading("Abstract", level=1)
        for p_el in abstract.findall("{*}p"):
            p = doc.add_paragraph()
            _add_run(p, p_el)

    # ── Keywords ──────────────────────────────────────────────────────────────
    kwds = root.findall(".//{*}kwd")
    if kwds:
        kwd_para = doc.add_paragraph()
        r = kwd_para.add_run("Keywords: ")
        r.bold = True
        for i, k in enumerate(kwds):
            if i > 0:
                kwd_para.add_run(", ")
            _add_run(kwd_para, k)

    doc.add_page_break()

    # ── Body sections ─────────────────────────────────────────────────────────
    body_el = root.find("{*}body")
    if body_el is not None:
        for sec in body_el.findall("{*}sec"):
            _docx_sec(doc, sec, level=1)

    # ── Back matter ───────────────────────────────────────────────────────────
    back = root.find("{*}back")
    if back is not None:
        for sec in back.findall("{*}sec"):
            _docx_sec(doc, sec, level=1)

    return doc


def _docx_sec(doc, sec, level):
    title_el = sec.find("{*}title")
    title_txt = _plain(title_el) if title_el is not None else ""
    if title_txt.strip():
        doc.add_heading(title_txt.strip(), level=min(level, 4))

    for child in sec:
        ctag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if ctag == "title":
            continue
        elif ctag == "p":
            p = doc.add_paragraph()
            _add_run(p, child)
        elif ctag == "sec":
            _docx_sec(doc, child, level + 1)
        elif ctag == "table-wrap":
            _docx_table(doc, child)
        elif ctag == "fig":
            _docx_figure(doc, child)
        elif ctag == "ref-list":
            _docx_refs(doc, child)


def _docx_table(doc, tbl_wrap):
    label = _plain(tbl_wrap.find("{*}label"))
    cap   = tbl_wrap.find(".//{*}caption")
    title_el  = cap.find("{*}title") if cap is not None else None

    # Table caption — preserve italic/bold formatting
    p = doc.add_paragraph()
    p.add_run(label.strip()).bold = True
    if cap is not None:
        if title_el is not None:
            p.add_run(" ")
            _add_run(p, title_el)
        for p_el in cap.findall("{*}p"):
            p.add_run(" ")
            _add_run(p, p_el)

    tbl = tbl_wrap.find(".//{*}table")
    if tbl is None:
        return
    rows = tbl.findall(".//{*}tr")
    if not rows:
        return

    # Count columns properly, accounting for colspan
    max_cols = max(
        sum(int(c.get("colspan", 1)) for c in row)
        for row in rows
    )

    dtbl = doc.add_table(rows=len(rows), cols=max_cols)
    try:
        dtbl.style = "Table Grid"
    except Exception:
        pass
    _set_table_borders(dtbl)

    # rowspan_remaining: {col_index: rows_still_occupied}
    rowspan_remaining: dict[int, int] = {}

    for ri, row in enumerate(rows):
        ci = 0
        for cell in row:
            # Skip columns occupied by a rowspan from a previous row
            while ci in rowspan_remaining:
                ci += 1
            colspan = int(cell.get("colspan", 1))
            rowspan = int(cell.get("rowspan", 1))
            if ci < max_cols:
                dtbl.cell(ri, ci).text = _plain(cell)
            # Register this cell's rowspan so subsequent rows skip its columns
            if rowspan > 1:
                for sc in range(ci, ci + colspan):
                    rowspan_remaining[sc] = rowspan - 1
            ci += colspan

        # Decrement counters; remove columns that have been fully consumed
        rowspan_remaining = {k: v - 1 for k, v in rowspan_remaining.items() if v > 0}

    doc.add_paragraph()  # spacing after table


def _docx_figure(doc, fig):
    global _XML_DIR
    label   = _plain(fig.find("{*}label"))
    cap     = fig.find("{*}caption")
    title_el = cap.find("{*}title") if cap is not None else None

    graphic = fig.find("{*}graphic")
    src = ""
    if graphic is not None:
        src = graphic.get("{http://www.w3.org/1999/xlink}href") or ""

    is_missing = not src or src.startswith("[MISSING:") or src.startswith("PLACEHOLDER")
    if not is_missing:
        img_path = (_XML_DIR / src).resolve()
        if img_path.exists():
            try:
                if _PIL_AVAILABLE:
                    # Convert to RGB PNG in memory — handles RGBA, 16-bit, CMYK, etc.
                    pil_img = _PILImage.open(img_path).convert("RGB")
                    buf = _io.BytesIO()
                    pil_img.save(buf, format="PNG")
                    buf.seek(0)
                    doc.add_picture(buf, width=Inches(5.5))
                else:
                    doc.add_picture(str(img_path), width=Inches(5.5))
            except Exception as e:
                doc.add_paragraph(f"[Figure error: {type(e).__name__}: {e} — {img_path}]")
        else:
            doc.add_paragraph(f"[Figure not found: {src}]")
    else:
        doc.add_paragraph(f"[{src or 'No image provided'}]")

    # Figure caption — preserve italic/bold formatting
    p = doc.add_paragraph()
    p.add_run(label.strip()).bold = True
    if cap is not None:
        if title_el is not None:
            p.add_run(" ")
            _add_run(p, title_el)
        for p_el in cap.findall("{*}p"):
            p.add_run(" ")
            _add_run(p, p_el)
    doc.add_paragraph()  # spacing after figure


def _docx_refs(doc, ref_list):
    for ref in ref_list.findall("{*}ref"):
        ref_id = ref.get("id", "")
        ec = ref.find(".//{*}element-citation") or ref.find(".//{*}mixed-citation")
        if ec is None:
            continue
        authors = []
        for pg in ec.findall("{*}person-group"):
            for name in pg.findall("{*}name"):
                s = _plain(name.find("{*}surname"))
                g = _plain(name.find("{*}given-names"))
                authors.append(f"{s} {g[0]}" if g else s)
            for collab in pg.findall("{*}collab"):
                authors.append(_plain(collab))
        title  = _plain(ec.find("{*}article-title"))
        source = _plain(ec.find("{*}source"))
        year   = _plain(ec.find("{*}year"))
        vol    = _plain(ec.find("{*}volume"))
        fp     = _plain(ec.find("{*}fpage"))
        lp     = _plain(ec.find("{*}lpage"))

        bib = f"{', '.join(authors)} ({year}). {title}. {source}"
        if vol:
            bib += f", {vol}"
        if fp:
            bib += f":{fp}–{lp}" if lp else f":{fp}"
        bib = bib.rstrip(". ") + "."

        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.add_run(bib)
        if ref_id:
            _add_bookmark(p, ref_id)


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xml_file")
    ap.add_argument("-o", "--output", default=None)
    ap.add_argument("--docx", action="store_true", help="Output Word document instead of HTML")
    args = ap.parse_args()

    xml_path = Path(args.xml_file)
    if not xml_path.exists():
        sys.exit(f"File not found: {xml_path}")

    if args.docx:
        if not _DOCX_AVAILABLE:
            sys.exit(
                "python-docx not installed. Run with:\n"
                "  uv run --with python-docx render_genome_note.py <xml> --docx"
            )
        global _XML_DIR
        _XML_DIR = xml_path.parent
        out_path = Path(args.output) if args.output else xml_path.with_suffix(".docx")
        doc = render_docx(xml_path)
        doc.save(str(out_path))
        print(f"Written: {out_path}")
        return

    out_path = Path(args.output) if args.output else xml_path.with_suffix(".html")
    is_review = "_review" in xml_path.name

    html = render(xml_path)

    if is_review:
        banner = (
            '<div class="review-banner">'
            '<b>REVIEW MODE</b> &mdash; '
            'Values highlighted in <strong>blue</strong> were pulled from the database. '
            'Unhighlighted text is static template content, manually entered, or marked [MISSING:].'
            '</div>'
        )
        html = html.replace('<div class="article-header">', f'{banner}\n    <div class="article-header">', 1)

    out_path.write_text(html, encoding="utf-8")
    print(f"Written: {out_path}")
    print(f"Open in browser: file://{out_path.resolve()}")


if __name__ == "__main__":
    main()
